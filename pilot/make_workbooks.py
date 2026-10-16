"""Two annotator workbooks (same 210 rows, different order; repeats >=60 rows from their originals)."""
import random, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
P = pd.read_csv('pilot/pilot_key_DO_NOT_SEND.csv').fillna('')
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); N = Font(name=F, size=11); B = Font(name=F, bold=True)
HF = PatternFill('solid', fgColor='1F3864'); Y = PatternFill('solid', fgColor='FFF2CC'); thin = Side(style='thin', color='BFBFBF')
CHOICES = [('Y', 'Yes - split here'), ('N', 'No - keep as one word here (the pieces are fine)'),
           ('E', 'Either is defensible'), ('X', 'The division is wrong (pieces are not right / cannot give this form)'),
           ('U', 'Unsure')]

def order(seed):
    rng = random.Random(seed)
    while True:
        ids = list(P.id); rng.shuffle(ids); pos = {i: k for k, i in enumerate(ids)}
        if all(abs(pos[r.id] - pos[r.repeat_of]) >= 60 for r in P[P.repeat_of != ''].itertuples()): return ids

def build(name, seed):
    wb = Workbook(); ws = wb.active; ws.title = 'Instructions'
    lines = [('Boundary judgements - pilot (Paper 2)', B), ('', N),
             ('Each row shows one written Sanskrit word and ONE point where it could be divided, marked with |.', N),
             ('The two pieces on either side of that point are shown in dictionary form.', N),
             ('Question: should a segmenter divide the word at this point?', B), ('', N)]
    for c, t in CHOICES: lines.append((f'   {c}  =  {t}', N))
    lines += [('', N), ('Only judge the marked point. Other possible divisions in the same word do not matter.', N),
              ('Where sounds merge at the join (e.g. tava + anyatra -> tavānyatra), the | is only approximate: the Pieces column is what you are judging.', N),
              ('E is a real answer, not a cop-out: use it when both choices are grammatically defensible.', N),
              ('U is better than a forced guess. Notes are welcome whenever the choice feels wrong.', N),
              ('Some rows may look familiar; please answer each one on its own.', N), ('', N),
              ('What to fill in: ONLY the yellow Answer column (pick from the dropdown) and, optionally, Note.', B),
              ('Work in the "Items" sheet. Save the file and send it back as it is.', N), ('', N),
              ('Example (not part of the task):', B)]
    for t, f in lines: ws.append([t]); ws.cell(ws.max_row, 1).font = f
    ex_h = ['Word', 'Split point', 'Pieces at this point', 'Answer', 'Note']
    ws.append(ex_h); r0 = ws.max_row
    for k in range(1, 6): ws.cell(r0, k).font = H; ws.cell(r0, k).fill = HF
    ws.append(['rāmeṇoktam', 'rāmeṇ | oktam', 'rāmeṇa + uktam', 'Y', 'two separate words'])
    for k in range(1, 6): ws.cell(ws.max_row, k).font = N
    ws.cell(ws.max_row, 4).fill = Y
    ws.column_dimensions['A'].width = 60
    for col in 'BCDE': ws.column_dimensions[col].width = 22
    it = wb.create_sheet('Items')
    hdr = ['#', 'ID', 'Word (IAST)', 'Word (Devanagari)', 'Split point (IAST)', 'Split point (Devanagari)',
           'Pieces at this point (IAST)', 'Pieces (Devanagari)', 'Answer', 'Note']
    it.append(hdr)
    for k in range(1, len(hdr) + 1):
        c = it.cell(1, k); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
    byid = P.set_index('id')
    for n, i in enumerate(order(seed), 1):
        r = byid.loc[i]
        it.append([n, i, r.word, r.word_deva, r.marked_iast, r.marked_deva, r.pieces_iast, r.pieces_deva, None, None])
        for k in range(1, len(hdr) + 1):
            c = it.cell(n + 1, k); c.font = N; c.border = Border(bottom=thin)
        it.cell(n + 1, 9).fill = Y
    dv = DataValidation(type='list', formula1='"Y,N,E,X,U"', allow_blank=True, showErrorMessage=True,
                        errorTitle='Use the list', error='Please choose Y, N, E, X or U.')
    it.add_data_validation(dv); dv.add(f'I2:I{it.max_row}')
    for col, w in zip('ABCDEFGHIJ', (5, 8, 34, 30, 36, 32, 28, 26, 9, 36)): it.column_dimensions[col].width = w
    it.freeze_panes = 'C2'; wb.active = 1
    wb.save(name)

build('pilot/pilot_annotator_A.xlsx', 101); build('pilot/pilot_annotator_B.xlsx', 202)
print('ok')
